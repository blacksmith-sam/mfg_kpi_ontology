"""Generate synthetic weekly reports from a fictional contract manufacturer
("ACME CM") so the adapter, history and propose steps can be tested.

  .venv/bin/python scripts/make_sample_cm_reports.py   -> data/inbox/ACME_weekly_2026-W36.xlsx ... W41 (+ W41 as .md)

The layout mimics a typical CM pack: a Summary sheet with Actual vs CM Target,
a Production matrix by stage, a Downtime event log, Quality and Traceability sheets.
"""
import os
import random
import sys

import openpyxl
from openpyxl.styles import Font

OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "data", "inbox")
WEEKS = [36, 37, 38, 39, 40, 41]


def week_data(t, rng):
    """t = 0..5, the line improves over time; W41 is close to the model baseline."""
    imp = t / 5
    n = lambda mu, sd: max(0, rng.gauss(mu, sd))
    started = int(4300 + 900 * imp + rng.gauss(0, 60))
    build_fpy = min(0.99, 0.925 + 0.031 * imp + rng.gauss(0, 0.004))
    test_fpy = min(0.99, 0.935 + 0.024 * imp + rng.gauss(0, 0.004))
    shipped = int(started * 0.995)
    d = {
        "units_shipped": shipped, "po_due": 120, "po_otif": int(98 + 11 * imp + rng.gauss(0, 2)),
        "dock_to_stock": round(n(9 - 2.5 * imp, 0.4), 1), "shortage_stops": int(n(6 - 3 * imp, 1)),
        "planned_hrs": 80, "wip": int(n(520 - 110 * imp, 15)), "queue_h": round(n(2.6 - 0.8 * imp, 0.15), 2),
        "kits_issued": 260, "kits_complete": int(244 + 8 * imp + rng.gauss(0, 1.5)),
        "loc_counted": 400, "loc_accurate": int(386 + 7 * imp + rng.gauss(0, 1.5)),
        "parts_received": int(n(40000 + 8000 * imp, 800)), "op_hours": int(n(560 - 80 * imp, 10)),
        "dl_sc": int(n(330 - 110 * imp, 10)), "dl_asm": int(n(980 - 280 * imp, 20)),
        "material_cost": int(started * 50 + rng.gauss(0, 3000)), "overhead": 95000,
        "rework_h": int(n(110 - 46 * imp, 6)), "rma": int(n(6 - 3 * imp, 1)), "scrap_cost": int(n(7200 - 3100 * imp, 300)),
        "build": {"started": started, "good": int(started * build_fpy), "cycle": round(n(58 - 6 * imp, 0.8), 1), "planned": 4800},
        "manual": {"started": 900, "good": int(900 * min(0.99, 0.955 + 0.015 * imp + rng.gauss(0, 0.003))),
                   "cycle": round(n(335 - 17 * imp, 4), 0)},
        "test": {"started": int(started * 1.085), "good": int(started * 1.085 * test_fpy), "cycle": 38,
                 "avail": round(min(0.99, 0.90 + 0.03 * imp + rng.gauss(0, 0.006)), 3), "stations": 1, "ndf": int(n(70 - 29 * imp, 4))},
        "pack": {"started": int(started * 1.02), "good": int(started * 1.02 * 0.996), "cycle": round(n(50 - 2 * imp, 0.5), 1), "planned": 4800},
        "incoming_rejects": int(n(70 - 32 * imp, 5)), "escapes": int(n(8 - 4 * imp, 1)),
        "cpk": round(n(1.22 + 0.19 * imp, 0.04), 2), "grr": round(n(16 - 4 * imp, 0.5), 1),
        "fast_ops": 10800, "fast_ok": int(10800 * min(0.9995, 0.993 + 0.0033 * imp + rng.gauss(0, 0.0004))),
        "labels": int(started * 1.02), "label_err": int(n(7 - 4 * imp, 1)),
        "bypass": int(n(14 - 8 * imp, 2)), "caught": int(n(260 - 30 * imp, 10)), "escaped_gate": int(n(16 - 7 * imp, 1.5)),
        "sw_audits": 50, "sw_ok": int(n(41 + 5 * imp, 1)), "esd_checks": 1440, "esd_fail": int(n(7 - 4 * imp, 1)),
        "genealogy": int(shipped * min(0.9995, 0.975 + 0.017 * imp + rng.gauss(0, 0.002))),
        "fields_req": shipped * 30, "drill_min": round(n(9 - 5 * imp, 0.5), 1), "integrity": int(rng.random() < 0.15),
        "dupes": int(n(3 - 2 * imp, 0.6)),
    }
    d["fields_cap"] = int(d["fields_req"] * min(0.9995, 0.986 + 0.007 * imp + rng.gauss(0, 0.001)))
    # downtime log: operator assists + breakdowns + other
    log = []
    for stage, total, n_assist, n_brk in (("Build", n(820 - 210 * imp, 40), int(n(150 - 65 * imp, 8)), int(n(10 - 4 * imp, 1))),
                                          ("Pack", n(420 - 120 * imp, 25), int(n(18 - 6 * imp, 3)), int(n(3 - 1 * imp, 0.7)))):
        used = 0
        for _ in range(n_assist):
            mins = rng.choice([1, 1, 2, 2, 3, 4])
            log.append((stage, mins, "Operator assist", rng.choice(["Part jam cleared", "Gripper reset", "Vision re-teach", "Feeder refill"])))
            used += mins
        for _ in range(n_brk):
            mins = int(n(85, 25))
            log.append((stage, mins, "Breakdown", rng.choice(["Servo fault", "Robot collision stop", "Screwdriver fault", "Conveyor jam"])))
            used += mins
        if total > used:
            log.append((stage, int(total - used), "Changeover / waiting", "Product changeover and material wait"))
    rng.shuffle(log)
    d["downtime_log"] = log
    return d


def write_xlsx(path, week_label, d):
    wb = openpyxl.Workbook()
    b = Font(bold=True)
    ws = wb.active
    ws.title = "Summary"
    ws.append(["ACME Contract Manufacturing - Weekly Operations Report"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Reporting week", week_label])
    ws.append([])
    ws.append(["Metric", "Actual", "CM Target", "Unit"])
    for c in ws[4]:
        c.font = b
    rows = [
        ("Units shipped", d["units_shipped"], None, "units"), ("PO lines due", d["po_due"], None, "lines"),
        ("PO lines delivered OTIF", d["po_otif"], None, "lines"),
        ("OTIF %", f"{d['po_otif'] / d['po_due']:.1%}", "95%", "%"),
        ("Dock-to-stock (avg hrs)", d["dock_to_stock"], 4, "h"),
        ("Line stops - material shortage", d["shortage_stops"], 0, "events"),
        ("Planned line hours", d["planned_hrs"], None, "h"), ("WIP at week end", d["wip"], None, "units"),
        ("Avg gate queue (h)", d["queue_h"], None, "h"), ("Kits issued", d["kits_issued"], None, "kits"),
        ("Kits complete at line", d["kits_complete"], None, "kits"), ("Cycle-count locations", d["loc_counted"], None, "loc"),
        ("Locations accurate", d["loc_accurate"], None, "loc"), ("Parts received", d["parts_received"], None, "units"),
        ("Operator hours - manual stations", d["op_hours"], None, "h"),
        ("Direct labor hours - supply chain", d["dl_sc"], None, "h"), ("Direct labor hours - assembly", d["dl_asm"], None, "h"),
        ("Material cost", f"${d['material_cost']:,}", None, "USD"), ("CM fee & overhead", f"${d['overhead']:,}", None, "USD"),
        ("Rework hours", d["rework_h"], 40, "h"), ("Customer returns (RMA)", d["rma"], 2, "units"),
        ("Scrap cost", f"${d['scrap_cost']:,}", "$3,000", "USD"),
    ]
    for r in rows:
        ws.append(list(r))
    ws.column_dimensions["A"].width = 36

    ws = wb.create_sheet("Production")
    ws.append(["Metric", "Build (auto)", "Manual stations", "Test", "Pack"])
    for c in ws[1]:
        c.font = b
    B, M, T, P = d["build"], d["manual"], d["test"], d["pack"]
    ws.append(["Units started", B["started"], M["started"], T["started"], P["started"]])
    ws.append(["First pass good", B["good"], M["good"], T["good"], P["good"]])
    ws.append(["Avg cycle time (s)", B["cycle"], M["cycle"], T["cycle"], P["cycle"]])
    ws.append(["Planned minutes", B["planned"], None, None, P["planned"]])
    ws.append(["Station availability", None, None, f"{T['avail']:.1%}", None])
    ws.append(["Stations in use", None, None, T["stations"], None])
    ws.append(["No defect found (NDF)", None, None, T["ndf"], None])
    ws.column_dimensions["A"].width = 26

    ws = wb.create_sheet("Downtime")
    ws.append(["Event", "Stage", "Minutes", "Category", "Reason"])
    for c in ws[1]:
        c.font = b
    for i, (stage, mins, cat, reason) in enumerate(d["downtime_log"], 1):
        ws.append([i, stage, mins, cat, reason])

    ws = wb.create_sheet("Quality")
    ws.append(["Metric", "Actual", "CM Target"])
    for c in ws[1]:
        c.font = b
    ppm = d["escapes"] / d["units_shipped"] * 1e6
    for r in [("Incoming rejects (units)", d["incoming_rejects"], None), ("Test escapes (units)", d["escapes"], None),
              ("Escape PPM", round(ppm), 1000), ("First pass yield - Test", f"{T['good'] / T['started']:.1%}", "95%"),
              ("Cpk critical position (datum A)", d["cpk"], 1.33), ("Gauge R&R % of tolerance", f"{d['grr']}%", "10%"),
              ("Fastening operations", d["fast_ops"], None), ("Fastening OK first time", d["fast_ok"], None),
              ("Label scans", d["labels"], None), ("Label errors", d["label_err"], None),
              ("Gate bypasses", d["bypass"], 0), ("Defects caught at gates", d["caught"], None),
              ("Defects escaped gates", d["escaped_gate"], None), ("Standard work audits", d["sw_audits"], None),
              ("Standard work audits compliant", d["sw_ok"], None), ("ESD checks", d["esd_checks"], None),
              ("ESD check fails", d["esd_fail"], None)]:
        ws.append(list(r))
    ws.column_dimensions["A"].width = 34

    ws = wb.create_sheet("Traceability")
    ws.append(["Metric", "Actual"])
    for r in [("Units with full genealogy", d["genealogy"]), ("Trace fields required", d["fields_req"]),
              ("Trace fields captured", d["fields_cap"]), ("Trace drill retrieval (min)", d["drill_min"]),
              ("Record integrity violations", d["integrity"]), ("Duplicate serials", d["dupes"])]:
        ws.append(list(r))
    wb.save(path)


def write_md(path, xlsx_path):
    """Same report as Markdown tables, as if the CM had sent it as text/PDF."""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    lines = []
    for ws in wb.worksheets:
        lines.append(f"## {ws.title}\n")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            cells = ["" if c is None else str(c) for c in row]
            if ws.title == "Summary" and i < 3:
                if any(cells):
                    lines.append(" ".join(c for c in cells if c) if i == 0 else f"{cells[0]}: {cells[1]}")
                continue
            lines.append("| " + " | ".join(cells) + " |")
        lines.append("")
    with open(path, "w") as fh:
        fh.write("\n".join(lines))


def main():
    os.makedirs(OUT, exist_ok=True)
    rng = random.Random(41)
    for t, w in enumerate(WEEKS):
        p = os.path.join(OUT, f"ACME_weekly_2026-W{w}.xlsx")
        write_xlsx(p, f"2026-W{w}", week_data(t, rng))
        print("wrote", p)
    md = os.path.join(OUT, "ACME_weekly_2026-W41_text_version.md")
    write_md(md, os.path.join(OUT, "ACME_weekly_2026-W41.xlsx"))
    print("wrote", md)


if __name__ == "__main__":
    sys.exit(main())
