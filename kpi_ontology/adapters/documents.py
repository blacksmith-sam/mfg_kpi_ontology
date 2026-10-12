"""Turn any report file into one neutral `Document`: a list of 2-D grids
(sheets) plus the full text. All locators work on this, so the same mapping
rules work for Excel, CSV, PDF or plain-text reports."""
from __future__ import annotations

import csv
import json
import os
import re
from dataclasses import dataclass, field


@dataclass
class Sheet:
    name: str
    rows: list  # list[list[str | float | None]]

    def cell(self, r, c):
        if 0 <= r < len(self.rows) and 0 <= c < len(self.rows[r]):
            return self.rows[r][c]
        return None


@dataclass
class Document:
    path: str
    sheets: list
    text: str
    meta: dict = field(default_factory=dict)

    def sheet(self, name=None):
        if name is None:
            return self.sheets
        hits = [s for s in self.sheets if s.name.lower() == str(name).lower()]
        hits = hits or [s for s in self.sheets if str(name).lower() in s.name.lower()]
        # text/PDF reports have no named sheets: fall back to searching everything
        return hits or self.sheets


def _grid_text(sheets):
    return "\n".join("\t".join("" if c is None else str(c) for c in row) for s in sheets for row in s.rows)


def _split_line(line):
    """Split a text line into cells on tabs, pipes (Markdown tables) or 2+ spaces."""
    line = line.strip().strip("|")
    if "|" in line:  # keep empty cells so table columns stay aligned
        cells = [c.strip() for c in line.split("|")]
        return [] if all(set(c) <= set("-: ") for c in cells) else [c if c else None for c in cells]
    else:
        cells = [c.strip() for c in re.split(r"\t|\s{2,}", line)]
    # also split "Label: value" into two cells
    if len(cells) == 1 and re.match(r"^[^:]{2,80}:\s*\S", cells[0]):
        cells = [x.strip() for x in cells[0].split(":", 1)]
    return [c for c in cells if c != ""]


def _text_to_sheets(text, name):
    rows = [_split_line(ln) for ln in text.splitlines()]
    return [Sheet(name, [r for r in rows if r])]


def load_xlsx(path):
    import openpyxl  # optional dependency
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheets = [Sheet(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]
    wb.close()
    return Document(path, sheets, _grid_text(sheets), {"format": "xlsx"})


def load_csv(path):
    with open(path, newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.reader(fh))
    sheets = [Sheet(os.path.splitext(os.path.basename(path))[0], rows)]
    return Document(path, sheets, _grid_text(sheets), {"format": "csv"})


def load_pdf(path):
    from pypdf import PdfReader  # optional dependency
    reader = PdfReader(path)
    sheets, texts = [], []
    for i, page in enumerate(reader.pages, 1):
        t = page.extract_text(extraction_mode="layout") or ""
        texts.append(t)
        sheets += _text_to_sheets(t, f"page{i}")
    return Document(path, sheets, "\n".join(texts), {"format": "pdf", "pages": len(reader.pages)})


def load_text(path):
    with open(path, encoding="utf-8", errors="replace") as fh:
        t = fh.read()
    return Document(path, _text_to_sheets(t, "text"), t, {"format": "text"})


def load_json(path):
    with open(path) as fh:
        data = json.load(fh)
    rows = []

    def flat(prefix, x):
        if isinstance(x, dict):
            for k, v in x.items():
                flat(f"{prefix}.{k}" if prefix else k, v)
        elif isinstance(x, list):
            for i, v in enumerate(x):
                flat(f"{prefix}[{i}]", v)
        else:
            rows.append([prefix, x])
    flat("", data)
    sheets = [Sheet("json", rows)]
    return Document(path, sheets, _grid_text(sheets), {"format": "json"})


LOADERS = {".xlsx": load_xlsx, ".xlsm": load_xlsx, ".csv": load_csv, ".pdf": load_pdf,
           ".txt": load_text, ".md": load_text, ".json": load_json}


def load_document(path):
    ext = os.path.splitext(path)[1].lower()
    if ext not in LOADERS:
        raise ValueError(f"no loader for {ext!r} files (have: {', '.join(LOADERS)}). "
                         "Export the report to one of these, or register a loader in LOADERS.")
    return LOADERS[ext](path)
